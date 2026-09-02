"""audit.py — 「統一されているか」と「同じ型の反復になっていないか」の検査。

`validate()`（slides.py）は**収まっているか**を見る検査ゲート。本モジュールは
その上に載る**編集の検査**で、生成済みの pptx を実測して次の2つを見る。

1. 統一（KEEP）— 統一してよい7項目が全ページで揃っているか
     基本フォントの方向性／本文色／アクセント色／余白の感覚（グリッド・基準線）／
     罫線の太さ／注釈の扱い（小さい文字は muted）／アクセントの点数

2. 変化（VARY）— 全ページが同じテンプレートになっていないか
     見出しの位置／カラム数／数値の見せ方／図解の有無／写真の有無／余白の位置／
     本文の密度／要素の大きさ／読み順／主役となる要素 を **pptx から再計測**し、
     連続ページで同じ構成が続いていないか・資料全体で構成が偏っていないかを見る

宣言（`spreads.Signature`）ではなく**描かれた結果**を測る。宣言だけ変えて中身が
同じ、という取り違えを防ぐため。

    from audit import audit
    audit(prs)          # ERROR があれば SystemExit(1)
"""
from __future__ import annotations

from dataclasses import dataclass

from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.dml import MSO_FILL

import slides as sl
from slides import Violation, print_report
import editorial as ed

__all__ = ["audit", "audit_violations", "Policy", "measure_signature"]

TOL = 0.12          # グリッド・基準線への吸着許容差（cm）
ZONE_Y0, ZONE_Y1 = 3.0, 12.8


# ============================================================================
# ポリシー（統一項目の許容集合）
# ============================================================================
@dataclass
class Policy:
    """統一してよい項目の許容値。既定は editorial.py のトークン。"""
    fonts: tuple = (ed.TYPE["jp"], ed.TYPE["latin"], ed.TYPE["display"],
                    "Courier New", "Arial", "Inter", "Noto Sans JP")
    inks: tuple = (ed.INK, ed.INK_SUB, ed.INK_REV, ed.ACCENT)
    rule_widths: tuple = tuple(sorted(set(ed.RULE.values())))
    accent_max: int = 3          # 1画面のアクセント要素数の上限（焦点1〜2＋余地1）
    note_size: float = ed.NOTE_SIZE
    grid: ed.Grid = ed.GRID
    min_variety: float = 0.6     # 相異なる構成 / ページ数 の下限
    near_distance: int = 3       # これ未満しか軸が違わなければ「ほぼ同じ誌面」


# ============================================================================
# 実測（pptx → 記録）
# ============================================================================
@dataclass
class Item:
    kind: str                    # text / shape / line / picture / chart
    x: float
    y: float
    w: float
    h: float
    fill: object = None          # RGBColor or None
    line: object = None          # RGBColor or None
    line_w: float = None         # pt
    runs: tuple = ()             # ((text, size, bold, rgb, latin_font), ...)
    align: str = "left"          # left / center / right（図に合わせた配置の判別用）

    @property
    def area(self):
        return max(self.w, 0) * max(self.h, 0)

    @property
    def max_size(self):
        return max((r[1] for r in self.runs if r[1]), default=0.0)


def _rgb(color_fmt):
    try:
        if color_fmt is None or color_fmt.type is None:
            return None
        return color_fmt.rgb
    except (AttributeError, TypeError, ValueError):
        return None


def _align_of(shape):
    """段落の横揃え。中央・右揃えのテキストは図や数値に合わせた配置とみなす。"""
    if not getattr(shape, "has_text_frame", False):
        return "left"
    for p in shape.text_frame.paragraphs:
        a = p.alignment
        if a is None:
            continue
        name = str(a).split(".")[-1].split(" ")[0].lower()
        if name in ("center", "right"):
            return name
    return "left"


def _runs_of(shape):
    out = []
    if not getattr(shape, "has_text_frame", False):
        return tuple(out)
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if not r.text.strip():
                continue
            size = r.font.size.pt if r.font.size is not None else None
            out.append((r.text, size, bool(r.font.bold),
                        _rgb(r.font.color), r.font.name))
    return tuple(out)


def _collect(slide):
    """スライド上の「自分で描いた要素」を実測する（テンプレ chrome は除く）。"""
    items = []
    for sh in slide.shapes:
        if sh.is_placeholder:
            continue          # 見出し帯・ページ番号はテンプレ契約（触らない）
        try:
            x, y = Emu(sh.left).cm, Emu(sh.top).cm
            w, h = Emu(sh.width).cm, Emu(sh.height).cm
        except (TypeError, ValueError):
            continue
        st = sh.shape_type
        if st == MSO_SHAPE_TYPE.PICTURE:
            kind = "picture"
        elif st == MSO_SHAPE_TYPE.CHART:
            kind = "chart"
        elif sh.__class__.__name__ == "Connector" or w < 0.06 or h < 0.06:
            kind = "line"
        elif getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
            kind = "text"
        else:
            kind = "shape"
        fill = line = None
        line_w = None
        try:
            if sh.fill.type == MSO_FILL.SOLID:
                fill = sh.fill.fore_color.rgb
        except (AttributeError, TypeError, ValueError, NotImplementedError):
            pass
        try:
            line = _rgb(sh.line.color)
            line_w = sh.line.width.pt if sh.line.width is not None else None
        except (AttributeError, TypeError, ValueError):
            pass
        items.append(Item(kind, x, y, w, h, fill, line, line_w, _runs_of(sh),
                          _align_of(sh)))
    return items


def _thin(it):
    """罫・バー・ひげのような「細い要素」か（図解の骨格になるもの）。"""
    if it.kind == "line":
        return True
    return min(it.w, it.h) <= 0.32 and max(it.w, it.h) >= 0.5


def _regions(items, *, gap=1.2, min_items=3, min_area=6.0):
    """細い要素が密集している範囲＝図解の領域を求める。

    図解の内部（線に合わせたラベル・バー・節点）は段組みの起点や基準線に乗らない
    のが正しいので、余白の検査からはこの領域を除く。
    """
    boxes = [[it.x, it.y, it.x + it.w, it.y + it.h, 1]
             for it in items if _thin(it)]
    merged = True
    while merged:
        merged = False
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                a, b = boxes[i], boxes[j]
                if (a[0] - gap <= b[2] and b[0] - gap <= a[2]
                        and a[1] - gap <= b[3] and b[1] - gap <= a[3]):
                    boxes[i] = [min(a[0], b[0]), min(a[1], b[1]),
                                max(a[2], b[2]), max(a[3], b[3]), a[4] + b[4]]
                    boxes.pop(j)
                    merged = True
                    break
            if merged:
                break
    return [b for b in boxes
            if b[4] >= min_items and (b[2] - b[0]) * (b[3] - b[1]) >= min_area]


def _slide_regions(slide, items):
    """そのスライドの図版領域。

    `ediagrams` が領域を申告しているデッキでは申告を正とする（申告が無いページは
    図版が無いページ）。申告が一切無いデッキだけ、細い要素の密集から推定する。
    """
    declared = ed.figure_areas(slide)
    if declared:
        return [list(r) + [99] for r in declared]
    if ed.has_figures(slide):
        return []
    return _regions(items)


def _inside(regions, it, pad=0.2):
    cx, cy = it.x + it.w / 2, it.y + it.h / 2
    return any(r[0] - pad <= cx <= r[2] + pad and r[1] - pad <= cy <= r[3] + pad
               for r in regions)


def _in_zone(it):
    cy = it.y + it.h / 2
    return ZONE_Y0 - 0.3 <= cy <= ZONE_Y1 + 0.3


# ============================================================================
# 実測シグネチャ（宣言ではなく描かれた結果から10軸を求める）
# ============================================================================
def _bucket_x(x, w):
    cx = x + w / 2
    third = ed.FIELD["w"] / 3
    if cx < ed.FIELD["x"] + third:
        return "left"
    if cx < ed.FIELD["x"] + third * 2:
        return "center"
    return "right"


def _bucket_y(y):
    third = ed.FIELD["h"] / 3
    if y < ed.FIELD["y"] + third:
        return "top"
    if y < ed.FIELD["y"] + third * 2:
        return "mid"
    return "bottom"


def _columns(items):
    """本文・見出しの左端をまとめて段の数を数える（溝で分かれた列＝段）。"""
    xs = sorted({round(it.x, 1) for it in items
                 if it.kind in ("text", "shape") and it.w >= 1.2})
    cols, last = [], None
    for x in xs:
        if last is None or x - last > 0.9:
            cols.append(x)
        last = x
    return cols


def measure_signature(slide) -> ed.Signature:
    """スライドを実測して 10 軸の Signature を作る。"""
    items = [it for it in _collect(slide) if _in_zone(it)]
    if not items:
        return ed.Signature(head_pos="none", columns="0", hero="none")

    texts = [it for it in items if it.kind == "text"]
    lines_ = [it for it in items if it.kind == "line"]
    pics = [it for it in items if it.kind == "picture"]
    charts = [it for it in items if it.kind == "chart"]
    regions = _slide_regions(slide, items)

    # 見出し＝最大の文字サイズを持つテキスト。位置は「どこから始まるか」で見る。
    head = max((it for it in texts if not _inside(regions, it)),
               key=lambda it: (it.max_size, it.area), default=None)
    head_pos = (f"{_bucket_x(head.x, 0)}-{_bucket_y(head.y)}" if head else "none")

    # 段＝図版をひとつの塊として数えた縦の列
    outside = [it for it in items if not _inside(regions, it)]
    cols = _columns(outside) + [round(r[0], 1) for r in regions]
    cols = sorted(cols)
    merged, last = [], None
    for c in cols:
        if last is None or c - last > 0.9:
            merged.append(c)
        last = c
    hero_box = max(
        [(r[0], r[1], r[2], r[3]) for r in regions]
        + [(it.x, it.y, it.x + it.w, it.y + it.h) for it in items],
        key=lambda b: (b[2] - b[0]) * (b[3] - b[1]))
    hcx = (hero_box[0] + hero_box[2]) / 2
    mid = ed.FIELD["x"] + ed.FIELD["w"] / 2
    side = ("L" if hcx < mid - ed.FIELD["w"] * 0.08 else
            "R" if hcx > mid + ed.FIELD["w"] * 0.08 else "C")
    columns = f"{len(merged)}:{side}"

    field_area = ed.FIELD["w"] * ed.FIELD["h"]
    diagram = bool(regions)
    photo = bool(pics)

    # 数値の見せ方
    big_num = any(r[1] and r[1] >= 24 and any(ch.isdigit() for ch in r[0])
                  for it in texts for r in it.runs)
    any_num = any(any(ch.isdigit() for ch in r[0])
                  for it in texts for r in it.runs)
    number = ("hero" if big_num else
              "chart" if charts else
              "inline" if any_num else "none")

    # 本文の密度＝文字数（面積ではなく実際の読む量で測る）
    chars = sum(len(r[0]) for it in texts for r in it.runs)
    density = "dense" if chars > 420 else ("sparse" if chars < 180 else "normal")

    # 主役＝面積最大の要素（種別で読み替える）
    hero_item = max(items, key=lambda it: it.area)
    if pics and max(pics, key=lambda i: i.area).area >= hero_item.area * 0.8:
        hero = "photo"
    elif charts:
        hero = "chart"
    elif regions and max((r[2] - r[0]) * (r[3] - r[1]) for r in regions) > field_area / 4:
        hero = "diagram"
    elif big_num:
        hero = "number"
    else:
        hero = "text"

    # 余白の位置＝4象限のうち最も要素の少ない側
    void = _void_quadrant(items, field_area)

    # 読み順＝主要素の並び
    order = _reading_order(items, regions)

    # 要素の大きさの差＝最大の級数が本文級数の何倍か
    sizes = sorted(r[1] for it in texts for r in it.runs if r[1])
    med = sizes[len(sizes) // 2] if sizes else 0
    scale = ("contrast" if med and max(sizes) / med >= 1.6 else "even")

    return ed.Signature(head_pos=head_pos, columns=columns, number=number,
                        diagram=diagram, photo=photo, void=void,
                        density=density, hero=hero, order=order, scale=scale)


def _void_quadrant(items, field_area):
    cx = ed.FIELD["x"] + ed.FIELD["w"] / 2
    cy = ed.FIELD["y"] + ed.FIELD["h"] / 2
    quad = {"top-left": 0.0, "top-right": 0.0,
            "bottom-left": 0.0, "bottom-right": 0.0}
    for it in items:
        icx, icy = it.x + it.w / 2, it.y + it.h / 2
        key = ("top" if icy < cy else "bottom") + ("-left" if icx < cx else "-right")
        quad[key] += it.area
    return min(quad, key=lambda k: quad[k])


def _reading_order(items, regions=()):
    """読み順。上から下に読ませるか、左右に渡らせるか、折り返させるか。"""
    big = sorted((it for it in items if it.area > 1.5), key=lambda it: it.y)
    if len(big) < 2:
        return "tb"
    if len(regions) >= 1 and len(big) >= 3:
        r = max(regions, key=lambda r: (r[2] - r[0]) * (r[3] - r[1]))
        if (r[2] - r[0]) > ed.FIELD["w"] * 0.8:
            return "radial" if (r[3] - r[1]) > ed.FIELD["h"] * 0.55 else "tb"
    mid = ed.FIELD["x"] + ed.FIELD["w"] / 2
    side = ["L" if (it.x + it.w / 2) < mid else "R" for it in big]
    flips = sum(1 for a, b in zip(side, side[1:]) if a != b)
    if flips >= 3:
        return "zig"
    xs = [it.x + it.w / 2 for it in big]
    ys = [it.y + it.h / 2 for it in big]
    spread_x = (max(xs) - min(xs)) / ed.FIELD["w"]
    spread_y = (max(ys) - min(ys)) / ed.FIELD["h"]
    return "lr" if spread_x >= spread_y else "tb"


# ============================================================================
# 検査
# ============================================================================
def _same_rgb(a, b):
    return a is not None and b is not None and str(a) == str(b)


def _check_consistency(page, items, pol, out, regions):
    accent_hits = 0
    for it in items:
        rect = (it.x, it.y, it.w, it.h)

        # 罫線の太さ
        if it.line_w is not None and it.line is not None:
            if not any(abs(it.line_w - v) < 0.06 for v in pol.rule_widths):
                out.append(Violation(
                    page, "WARN", "RULE_W", it.kind, rect,
                    f"罫線の太さが規定外 {it.line_w:.2f}pt "
                    f"（許容 {', '.join(f'{v:g}' for v in pol.rule_widths)}）"))
        # アクセントの点数
        if _same_rgb(it.fill, ed.ACCENT) or _same_rgb(it.line, ed.ACCENT):
            accent_hits += 1
        for text, size, _bold, rgb, latin in it.runs:
            if _same_rgb(rgb, ed.ACCENT):
                accent_hits += 1
            # 本文色・アクセント色
            if rgb is not None and not any(_same_rgb(rgb, c) for c in pol.inks):
                out.append(Violation(
                    page, "ERROR", "INK", it.kind, rect,
                    f"本文色・アクセント色の外の文字色 #{rgb} 「{text[:20]}」"))
            # 基本フォントの方向性
            if latin and latin not in pol.fonts:
                out.append(Violation(
                    page, "WARN", "FONT", it.kind, rect,
                    f"規定外の書体 {latin} 「{text[:20]}」"))
            # 注釈の扱い（小さい文字は必ず muted）
            if size and size <= pol.note_size + 0.3:
                if rgb is not None and not _same_rgb(rgb, ed.INK_SUB) \
                        and not _same_rgb(rgb, ed.INK_REV):
                    out.append(Violation(
                        page, "WARN", "NOTE_STYLE", it.kind, rect,
                        f"注釈サイズ({size:g}pt)の文字が補助色でない #{rgb} "
                        f"「{text[:20]}」"))
        # 余白の感覚（列の起点・基準線への吸着）
        # 図解の内部と注釈は対象外（線・図に合わせる位置合わせが正しいため）
        note_like = it.max_size and it.max_size <= pol.note_size + 0.3
        if (it.kind in ("text", "shape") and it.w >= 1.0
                and it.align == "left"
                and not note_like and not _inside(regions, it)):
            if not _on_column(it.x, pol.grid):
                out.append(Violation(
                    page, "WARN", "GRID_X", it.kind, rect,
                    f"左端が段組みの起点に乗っていない x={it.x:.2f}cm"))
            if not _on_baseline(it.y):
                out.append(Violation(
                    page, "WARN", "GRID_Y", it.kind, rect,
                    f"上端が基準線に乗っていない y={it.y:.2f}cm"))
    if accent_hits > pol.accent_max:
        out.append(Violation(
            page, "WARN", "ACCENT_MANY", "slide",
            (ed.FIELD["x"], ed.FIELD["y"], ed.FIELD["w"], ed.FIELD["h"]),
            f"アクセントが{accent_hits}箇所（焦点は1〜2箇所に絞る）"))


def _on_column(x, grid):
    for i in range(grid.n):
        if abs(x - grid.span(i, 1)[0]) < TOL:
            return True
    return abs(x - ed.FIELD["x"]) < TOL


def _on_baseline(y):
    n = round((y - ed.FIELD["y"]) / ed.BASE)
    return abs(ed.baseline(int(n)) - y) < TOL


def _check_symmetry(page, items, out):
    """左右対称・等間隔・同サイズだけで組まれた誌面を拾う。"""
    body = [it for it in items if it.kind in ("text", "shape") and it.w >= 1.2]
    if len(body) < 3:
        return
    left = min(it.x for it in body) - ed.FIELD["x"]
    right = (ed.FIELD["x"] + ed.FIELD["w"]) - max(it.x + it.w for it in body)
    widths = sorted({round(it.w, 1) for it in body})
    sizes = [r[1] for it in body for r in it.runs if r[1]]
    even_w = len(widths) <= 1
    even_size = bool(sizes) and (max(sizes) / min(sizes) < 1.4)
    if abs(left - right) < 0.15 and even_w and even_size:
        out.append(Violation(
            page, "WARN", "SYMMETRY", "slide",
            (ed.FIELD["x"], ed.FIELD["y"], ed.FIELD["w"], ed.FIELD["h"]),
            "左右対称・等幅・同サイズだけの構成（余白の偏り／大きさの差を作る）"))


def audit_violations(prs, *, policy=None):
    """統一・変化の検査結果（Violation リスト）を返す。"""
    pol = policy or Policy()
    out: list[Violation] = []
    body_layout = sl._body_layout_name(prs)

    sigs: list[tuple[int, ed.Signature]] = []
    for i, s in enumerate(prs.slides):
        if s.slide_layout.name != body_layout:
            continue      # 表紙・章扉・裏表紙はテンプレ契約なので対象外
        page = i + 1
        items = [it for it in _collect(s) if _in_zone(it)]
        if not items:
            continue
        _check_consistency(page, items, pol, out, _slide_regions(s, items))
        _check_symmetry(page, items, out)
        sigs.append((page, measure_signature(s)))

    # ---- 変化（同じ誌面の反復）
    for (p0, a), (p1, b) in zip(sigs, sigs[1:]):
        d = a.distance(b)
        rect = (ed.FIELD["x"], ed.FIELD["y"], ed.FIELD["w"], ed.FIELD["h"])
        if d == 0:
            out.append(Violation(
                p1, "ERROR", "SAME_SPREAD", "slide", rect,
                f"P{p0} と構成が完全に同一（10軸すべて一致）"))
        elif d < pol.near_distance:
            out.append(Violation(
                p1, "WARN", "NEAR_SPREAD", "slide", rect,
                f"P{p0} と構成がほぼ同じ（違いは{d}軸のみ）"))

    if len(sigs) >= 4:
        uniq = len({s.key() for _, s in sigs})
        ratio = uniq / len(sigs)
        if ratio < pol.min_variety:
            out.append(Violation(
                sigs[-1][0], "WARN", "VARIETY_LOW", "deck",
                (ed.FIELD["x"], ed.FIELD["y"], ed.FIELD["w"], ed.FIELD["h"]),
                f"構成の種類が少ない {uniq}/{len(sigs)}（下限 {pol.min_variety:.0%}）"))
    return out


def audit(prs, *, policy=None, strict=True, verbose=True):
    """統一・変化の検査を実行してレポートする。ERROR があれば SystemExit(1)。"""
    import sys
    violations = audit_violations(prs, policy=policy)
    if verbose:
        _print_signatures(prs)
    n_err, _ = print_report(violations, title="編集検査（統一と変化）")
    if strict and n_err:
        sys.exit(1)
    return violations


def _print_signatures(prs):
    body_layout = sl._body_layout_name(prs)
    print("\n----- 実測した構成（10軸） -----")
    for i, s in enumerate(prs.slides):
        if s.slide_layout.name != body_layout:
            continue
        sig = measure_signature(s)
        print(f"P{i + 1:>2} 見出し={sig.head_pos:<12} 段={sig.columns:<3} "
              f"数値={sig.number:<6} 図={'○' if sig.diagram else '−'} "
              f"写真={'○' if sig.photo else '−'} 余白={sig.void:<12} "
              f"密度={sig.density:<6} 主役={sig.hero:<7} 読順={sig.order:<3} "
              f"大小={sig.scale}")
