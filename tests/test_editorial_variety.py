"""「同じテンプレートの反復になっていないか」を見る検査（audit.py）の回帰テスト。

audit は宣言ではなく **描かれた pptx** を実測する。そこで、
  - 同じ構成を並べたら反復として検出されること
  - 内容の形で構成が選ばれ、実測しても構成が散ること
  - 統一項目（本文色・アクセント色・罫線の太さ・注釈の色）の逸脱を拾うこと
を確認する。
"""
import pytest
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

import editorial as ed
import spreads as sp
from audit import audit_violations, measure_signature
from slides import L_BODY, configure_body, textbox, shape_box, connector


@pytest.fixture(autouse=True)
def _clean_figures():
    ed.reset_figures()
    yield
    ed.reset_figures()


def _codes(violations):
    return [v.code for v in violations]


# ---------------------------------------------------------------- 反復の検出
def _same_page(prs, n):
    """まったく同じ構成のページを n 枚作る（＝やってはいけない作り方）。"""
    for i in range(n):
        s = prs.slides.add_slide(prs.slide_layouts[L_BODY])
        configure_body(s, section_label=f"SECTION {i}", title=f"見出し {i}")
        x, w = ed.GRID.span(0, 6)
        ed.heading(s, x, ed.baseline(1), w, f"同じ場所に同じ見出し {i}")
        ed.body(s, x, ed.baseline(4), w, [f"同じ場所に同じ本文 {i}"] * 3)
        ed.rule(s, x, ed.baseline(12), w, weight="rule")


def test_identical_spreads_are_reported(prs):
    _same_page(prs, 3)
    codes = _codes(audit_violations(prs))
    assert "SAME_SPREAD" in codes or "NEAR_SPREAD" in codes


def test_measured_signature_ignores_declared_labels(prs):
    """宣言ではなく実測なので、同じ内容でも構成が違えば違う値になる。"""
    content = {"section": "S", "title": "T", "head": "H",
               "body": ["本文がここにある。"] * 2, "note": ["注記"]}
    s1, _ = sp.sp_lead_and_void(prs, content)
    s2, _ = sp.sp_two_voices(prs, dict(content, aside_blocks=[("観測", "補足")]))
    assert measure_signature(s1).key() != measure_signature(s2).key()


# ---------------------------------------------------------------- 構成の選択
def test_choose_follows_the_shape_of_the_content():
    assert sp.choose({"steps": [("A", ""), ("B", "")]}) == "process_band"
    assert sp.choose({"pairs": [("A", "B"), ("C", "D")]}) == "correspondence"
    assert sp.choose({"load": [("A", 1), ("B", 2)]}) == "load_report"
    assert sp.choose({"points": [("1月", 1), ("2月", 2)]}) == "change_note"
    assert sp.choose({"nodes": [("A", ""), ("B", "")]}) == "connection_map"
    assert sp.choose({"value": "87"}) == "figure_first"


def test_choose_avoids_repeating_the_recent_spread():
    text_only = {"body": ["文章だけのページ。" * 4]}
    first = sp.choose(text_only, [])
    second = sp.choose(text_only, [first])
    assert second != first


def test_render_all_produces_varied_spreads(prs):
    contents = [
        {"section": "A", "title": "1", "statement": "短い1文。"},
        {"section": "B", "title": "2", "value": "87", "unit": "%",
         "head": "解釈", "body": ["数値の読み方。"]},
        {"section": "C", "title": "3", "head": "工程",
         "steps": [("企画", "範囲"), ("設計", "前提"), ("制作", "本体")],
         "body": ["工程の説明。"]},
        {"section": "D", "title": "4", "head": "対応",
         "pairs": [("問い", "打ち手"), ("問い2", "打ち手2")]},
    ]
    used = sp.render_all(prs, contents)
    assert len({name for name, _ in used}) == len(contents)
    sigs = [measure_signature(s) for s in prs.slides
            if s.slide_layout.name == prs.slide_layouts[L_BODY].name]
    assert len({s.key() for s in sigs}) == len(sigs)
    assert not [v for v in audit_violations(prs) if v.severity == "ERROR"]


def test_render_all_stops_on_tone_violation(prs):
    with pytest.raises(ValueError):
        sp.render_all(prs, [{"section": "A", "title": "煽り",
                             "statement": "今すぐ買わないと手遅れになります"}])


# ---------------------------------------------------------------- 統一の逸脱
def test_off_palette_text_is_an_error(prs):
    s = prs.slides.add_slide(prs.slide_layouts[L_BODY])
    configure_body(s, section_label="S", title="T")
    textbox(s, 1.07, 5.0, 8.0, 1.0, ["規定外の色の文字"], size=11,
            color=RGBColor(0xCC, 0x00, 0x00))
    assert "INK" in _codes(audit_violations(prs))


def test_off_scale_rule_width_is_reported(prs):
    s = prs.slides.add_slide(prs.slide_layouts[L_BODY])
    configure_body(s, section_label="S", title="T")
    connector(s, 1.07, 6.0, 10.0, 6.0, color=ed.HAIR, width=3.4)
    assert "RULE_W" in _codes(audit_violations(prs))


def test_small_text_must_be_muted(prs):
    s = prs.slides.add_slide(prs.slide_layouts[L_BODY])
    configure_body(s, section_label="S", title="T")
    textbox(s, 1.07, 11.0, 8.0, 0.6, ["注釈のつもりの小さな文字"],
            size=ed.NOTE_SIZE, color=ed.INK)
    assert "NOTE_STYLE" in _codes(audit_violations(prs))


def test_symmetric_equal_layout_is_reported(prs):
    s = prs.slides.add_slide(prs.slide_layouts[L_BODY])
    configure_body(s, section_label="S", title="T")
    for i in range(3):                   # 等幅・等間隔で誌面を3等分する
        x = 1.07 + i * 7.92
        shape_box(s, MSO_SHAPE.RECTANGLE, x, 5.0, 7.42, 3.0,
                  text="同じ大きさの箱", fill=None, line=ed.HAIR,
                  line_w=ed.RULE["thin"], size=11, color=ed.INK)
    assert "SYMMETRY" in _codes(audit_violations(prs))
